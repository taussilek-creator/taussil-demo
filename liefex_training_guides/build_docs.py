from __future__ import annotations

import io
import os
import zipfile
import uuid
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

ORANGE = "E4520D"
ORANGE_DARK = "B83D08"
CHARCOAL = "20272D"
DARK = "11171B"
GRAY = "687078"
LIGHT = "F4F6F7"
LIGHT_ORANGE = "FFF1E9"
GREEN = "158A69"
LIGHT_GREEN = "EAF7F2"
RED = "C83C37"
LIGHT_RED = "FDEEEE"
BLUE = "2B6CB0"
WHITE = "FFFFFF"
FONT_NAME = "Almarai"
REGULAR_FONT = Path(os.environ.get("ALMARAI_REGULAR", "/usr/local/share/fonts/Almarai-Regular.ttf"))
BOLD_FONT = Path(os.environ.get("ALMARAI_BOLD", "/usr/local/share/fonts/Almarai-Bold.ttf"))


def ar(text: str) -> str:
    return get_display(arabic_reshaper.reshape(text))


def img_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = BOLD_FONT if bold and BOLD_FONT.exists() else REGULAR_FONT
    if not path.exists():
        path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    return ImageFont.truetype(str(path), size=size)


def rounded(draw: ImageDraw.ImageDraw, box, radius=24, fill=WHITE, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def txt(draw: ImageDraw.ImageDraw, xy, text: str, size=34, fill=CHARCOAL, bold=False, anchor="ra"):
    draw.text(xy, ar(text), font=img_font(size, bold), fill="#" + fill, anchor=anchor)


def brand(draw: ImageDraw.ImageDraw, x: int, y: int, size=48):
    f = img_font(size, True)
    draw.text((x, y), "Liefe", font=f, fill="#11171B", anchor="la")
    w = draw.textlength("Liefe", font=f)
    draw.text((x + w, y), "X", font=f, fill="#E4520D", anchor="la")


def phone_frame(draw, box, fill="#172027"):
    x1, y1, x2, y2 = box
    rounded(draw, box, radius=48, fill=fill, outline="#39434A", width=4)
    draw.rounded_rectangle((x1 + 80, y1 + 18, x2 - 80, y1 + 36), radius=10, fill="#090D10")


def draw_dashboard(path: Path):
    im = Image.new("RGB", (1400, 900), "#F5F6F7")
    d = ImageDraw.Draw(im)
    phone_frame(d, (430, 35, 970, 865))
    brand(d, 510, 90, 36)
    txt(d, (905, 98), "لوحة الإدارة", 28, WHITE, True)
    rounded(d, (500, 145, 900, 205), 18, "2B7775")
    txt(d, (860, 175), "الدعم", 24, WHITE, True)
    txt(d, (875, 250), "قائمة التوصيلات", 27, WHITE, True)
    txt(d, (670, 250), "متاح", 24, WHITE, True)
    d.ellipse((600, 232, 650, 282), fill="#FFFFFF")
    d.rounded_rectangle((625, 240, 690, 274), radius=16, fill="#2AA69A")
    for i, lab in enumerate(["الطلبات", "الأرباح", "الضبط", "السجل"]):
        x = 495 + i * 100
        rounded(d, (x, 785, x + 90, 830), 12, ORANGE if i == 0 else "293139")
        txt(d, (x + 80, 808), lab, 17, WHITE, True)
    # callouts
    for n, (cx, cy, label) in enumerate([(1030, 175, "الدعم"), (1030, 260, "حالة التوفر"), (1030, 800, "التبويبات"), (350, 120, "تسجيل الخروج")], 1):
        d.ellipse((cx - 26, cy - 26, cx + 26, cy + 26), fill="#E4520D")
        d.text((cx, cy), str(n), font=img_font(24, True), fill="white", anchor="mm")
        txt(d, (cx + 180 if cx > 700 else cx - 60, cy), label, 27, CHARCOAL, True, anchor="rm" if cx > 700 else "lm")
    im.save(path, quality=95)


def draw_incoming(path: Path, hybrid=False):
    im = Image.new("RGB", (1400, 900), "#F5F6F7")
    d = ImageDraw.Draw(im)
    rounded(d, (110, 125, 1290, 750), 36, WHITE, outline="#D9DEE1", width=3)
    txt(d, (1225, 185), "طلب هجين" if hybrid else "طلب أصيل", 44, CHARCOAL, True)
    rounded(d, (170, 245, 1230, 585), 28, "20272D", outline="#2B8D83", width=4)
    rounded(d, (1040, 300, 1175, 435), 24, WHITE)
    d.ellipse((1075, 330, 1140, 395), fill="#F0A43C")
    txt(d, (995, 315), "المتجر: Doner Haus", 30, WHITE, True)
    txt(d, (995, 375), "الكمية: 2", 27, WHITE)
    txt(d, (995, 430), "معرّف الطلب: 7DB8C9", 27, WHITE)
    if hybrid:
        txt(d, (995, 490), "لا يظهر عنوان العميل قبل القبول", 25, "E4520D", True)
    else:
        txt(d, (995, 490), "بيانات المتجر والعميل متاحة", 25, "66D6B7", True)
    rounded(d, (220, 485, 430, 550), 18, "C83C37")
    txt(d, (395, 518), "رفض", 25, WHITE, True)
    rounded(d, (470, 485, 680, 550), 18, "158A69")
    txt(d, (645, 518), "قبول", 25, WHITE, True)
    txt(d, (865, 535), "01:20", 34, "E4520D", True)
    # info ribbon
    rounded(d, (170, 635, 1230, 710), 18, LIGHT_ORANGE)
    txt(d, (1190, 673), "اتخذ القرار خلال المهلة. التجاهل يتحول إلى طلب فائت.", 27, ORANGE_DARK, True)
    im.save(path, quality=95)


def draw_flow(path: Path):
    im = Image.new("RGB", (1600, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    title = "رحلة الطلب داخل التطبيق"
    txt(d, (1510, 80), title, 48, CHARCOAL, True)
    steps = [
        ("1", "الجاهزية", "تفعيل متاح"),
        ("2", "القرار", "قبول أو رفض"),
        ("3", "المتجر", "استلام وتأمين"),
        ("4", "العميل", "ملاحة وتتبع"),
        ("5", "التسليم", "تحقق وتحصيل"),
        ("6", "الإغلاق", "تأكيد التسليم"),
    ]
    for i, (n, h, s) in enumerate(steps):
        x = 1470 - i * 245
        if i < len(steps) - 1:
            d.line((x - 120, 360, x - 220, 360), fill="#E4520D", width=8)
            d.polygon([(x - 220, 345), (x - 245, 360), (x - 220, 375)], fill="#E4520D")
        d.ellipse((x - 75, 280, x + 75, 430), fill="#E4520D")
        d.text((x, 355), n, font=img_font(52, True), fill="white", anchor="mm")
        txt(d, (x + 90, 500), h, 30, CHARCOAL, True, anchor="ra")
        txt(d, (x + 90, 550), s, 23, GRAY, False, anchor="ra")
    rounded(d, (120, 660, 1480, 790), 28, CHARCOAL)
    txt(d, (1420, 725), "لا تضغط «تم التسليم» قبل التسليم الفعلي والتحصيل عند الدفع النقدي.", 30, WHITE, True)
    im.save(path, quality=95)


def draw_pickup(path: Path):
    im = Image.new("RGB", (1400, 900), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1320, 75), "الاستلام المهني من المتجر", 45, CHARCOAL, True)
    # store counter
    d.rectangle((120, 500, 1280, 760), fill="#D9B38C")
    d.rectangle((120, 210, 1280, 500), fill="#FFF7EF")
    # merchant
    d.ellipse((910, 245, 1040, 375), fill="#D69A72")
    d.rectangle((900, 370, 1050, 585), fill="#30363B")
    # driver
    d.ellipse((330, 245, 460, 375), fill="#C88862")
    d.rectangle((315, 370, 475, 625), fill="#1C2328")
    d.rectangle((345, 430, 445, 610), fill="#E4520D")
    # bag on table
    rounded(d, (560, 390, 820, 590), 30, "E4520D", outline="#B83D08", width=5)
    d.rectangle((610, 355, 770, 405), fill="#20272D")
    txt(d, (760, 485), "LiefeX", 34, WHITE, True)
    # callout steps
    items = ["عرّف بنفسك", "أظهر معرّف الطلب", "افحص التغليف خارجياً", "ضع الطلب على سطح مرتفع", "أمّنه داخل الحقيبة"]
    for i, lab in enumerate(items):
        y = 110 + i * 95
        x = 70
        d.ellipse((x, y, x + 54, y + 54), fill="#E4520D")
        d.text((x + 27, y + 27), str(i + 1), font=img_font(24, True), fill="white", anchor="mm")
        txt(d, (x + 440, y + 27), lab, 25, CHARCOAL, True)
    im.save(path, quality=95)


def draw_bag(path: Path):
    im = Image.new("RGB", (1400, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1320, 70), "حماية الطلب داخل حقيبة التوصيل", 44, CHARCOAL, True)
    rounded(d, (430, 180, 970, 720), 40, "E4520D", outline="#B83D08", width=8)
    d.rectangle((470, 250, 930, 650), fill="#FFF8F2")
    d.line((700, 250, 700, 650), fill="#30363B", width=10)
    txt(d, (660, 320), "ساخن", 34, ORANGE_DARK, True)
    txt(d, (890, 320), "بارد", 34, BLUE, True)
    # cups / boxes
    for y in (390, 510):
        rounded(d, (505, y, 650, y + 90), 12, "F5B25D")
        rounded(d, (750, y, 895, y + 90), 12, "BFE3F7")
    callouts = [
        ("فصل الساخن عن البارد", 155, 245, GREEN),
        ("تثبيت السوائل والمشروبات", 155, 380, BLUE),
        ("منع الأغراض الشخصية", 155, 515, RED),
        ("تنظيف دوري للحقيبة", 155, 650, ORANGE),
    ]
    for label, x, y, color in callouts:
        rounded(d, (x, y, x + 230, y + 85), 20, color)
        txt(d, (x + 205, y + 42), label, 22, WHITE, True)
    im.save(path, quality=95)


def draw_cash(path: Path):
    im = Image.new("RGB", (1500, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1420, 70), "مسار الطلب النقدي والتسوية", 46, CHARCOAL, True)
    boxes = [
        (1170, "قيمة الطلب", "تظهر على الكرت"),
        (860, "التحصيل", "من العميل كاملاً"),
        (550, "نهاية النوبة", "تسليم الكاش للإدارة"),
        (240, "كشف التسوية", "خصم 5% من أرباح السائق"),
    ]
    for i, (x, h, s) in enumerate(boxes):
        rounded(d, (x - 125, 260, x + 125, 540), 28, WHITE, outline="#D8DEE2", width=3)
        d.ellipse((x - 40, 295, x + 40, 375), fill="#E4520D")
        d.text((x, 335), str(i + 1), font=img_font(34, True), fill="white", anchor="mm")
        txt(d, (x + 95, 425), h, 27, CHARCOAL, True)
        txt(d, (x + 95, 475), s, 21, GRAY)
        if i < 3:
            d.line((x - 140, 400, x - 180, 400), fill="#E4520D", width=8)
            d.polygon([(x - 180, 385), (x - 205, 400), (x - 180, 415)], fill="#E4520D")
    rounded(d, (170, 640, 1330, 760), 26, CHARCOAL)
    txt(d, (1270, 700), "الكاش المستحق ليس ربحاً للسائق؛ هو أموال طلبات تُسلّم كاملة للإدارة.", 28, WHITE, True)
    im.save(path, quality=95)


def draw_safety(path: Path):
    im = Image.new("RGB", (1400, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1320, 70), "السلامة أثناء العمل", 46, CHARCOAL, True)
    cards = [
        (70, "الخوذة إلزامية", "رأس محمي", GREEN),
        (400, "الهاتف ممنوع أثناء القيادة", "هاتف مشطوب", RED),
        (730, "لا مرافقين", "شخص واحد", ORANGE),
        (1060, "معدات المركبة كاملة", "فحص السلامة", BLUE),
    ]
    for x, h, s, c in cards:
        rounded(d, (x, 190, x + 270, 650), 32, WHITE, outline="#D9DEE1", width=3)
        d.ellipse((x + 70, 240, x + 200, 370), fill="#" + c)
        if "الهاتف" in h:
            d.rounded_rectangle((x + 112, 272, x + 158, 338), radius=8, outline="white", width=5)
            d.line((x + 90, 250, x + 180, 360), fill="white", width=9)
        elif "الخوذة" in h:
            d.arc((x + 96, 260, x + 174, 336), 190, 355, fill="white", width=9)
            d.line((x + 100, 335, x + 177, 335), fill="white", width=9)
        elif "مرافقين" in h:
            d.ellipse((x + 115, 275, x + 155, 315), fill="white")
            d.line((x + 135, 320, x + 135, 352), fill="white", width=9)
        else:
            d.line((x + 100, 325, x + 125, 350), fill="white", width=10)
            d.line((x + 125, 350, x + 180, 285), fill="white", width=10)
        txt(d, (x + 240, 440), h, 24, CHARCOAL, True)
        txt(d, (x + 240, 500), s, 20, GRAY)
    im.save(path, quality=95)


def draw_emergency(path: Path):
    im = Image.new("RGB", (1400, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1320, 70), "الاستجابة للحالات الطارئة", 46, CHARCOAL, True)
    d.polygon([(700, 170), (1010, 700), (390, 700)], fill="#E4520D")
    d.polygon([(700, 245), (935, 650), (465, 650)], fill="#FFF3EB")
    d.rectangle((674, 350, 726, 520), fill="#E4520D")
    d.ellipse((674, 555, 726, 607), fill="#E4520D")
    steps = ["أمّن نفسك والمركبة", "احفظ الطلب داخل الحقيبة", "اتصل بالإدارة", "اتصل بالجهات المختصة عند الحاجة", "لا تتصرف منفرداً في الطلب"]
    for i, s in enumerate(steps):
        y = 190 + i * 105
        d.ellipse((90, y, 145, y + 55), fill="#20272D")
        d.text((117, y + 27), str(i + 1), font=img_font(23, True), fill="white", anchor="mm")
        txt(d, (365, y + 27), s, 24, CHARCOAL, True)
    im.save(path, quality=95)


def draw_accreditation(path: Path):
    im = Image.new("RGB", (1500, 850), "#F5F6F7")
    d = ImageDraw.Draw(im)
    txt(d, (1420, 70), "مسار الاعتماد", 46, CHARCOAL, True)
    steps = ["حساب مكتمل", "مستندات", "دورة 6 ساعات", "اختبار نظري", "اختبار عملي", "تفعيل الحساب", "بطاقة سائق معتمد"]
    for i, s in enumerate(steps):
        x = 1390 - i * 205
        d.ellipse((x - 55, 260, x + 55, 370), fill="#E4520D" if i < 6 else "#158A69")
        d.text((x, 315), str(i + 1), font=img_font(30, True), fill="white", anchor="mm")
        txt(d, (x + 82, 450), s, 22, CHARCOAL, True)
        if i < 6:
            d.line((x - 60, 315, x - 135, 315), fill="#B9C0C4", width=7)
    rounded(d, (190, 620, 1310, 760), 28, CHARCOAL)
    txt(d, (1250, 690), "النجاح: 70% في النظري و70% في العملي، مع عدم ارتكاب خطأ رسوب مباشر.", 28, WHITE, True)
    im.save(path, quality=95)


def make_assets():
    funcs = {
        "dashboard.png": draw_dashboard,
        "original.png": lambda p: draw_incoming(p, False),
        "hybrid.png": lambda p: draw_incoming(p, True),
        "flow.png": draw_flow,
        "pickup.png": draw_pickup,
        "bag.png": draw_bag,
        "cash.png": draw_cash,
        "safety.png": draw_safety,
        "emergency.png": draw_emergency,
        "accreditation.png": draw_accreditation,
    }
    for name, fn in funcs.items():
        fn(ASSETS / name)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_rtl(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    bidi = tblPr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tblPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, size=None, bold=None, color=None, rtl=True):
    run.font.name = FONT_NAME
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), FONT_NAME)
    rtl_el = rPr.find(qn("w:rtl"))
    if rtl:
        if rtl_el is None:
            rtl_el = OxmlElement("w:rtl")
            rPr.append(rtl_el)
        rtl_el.set(qn("w:val"), "1")
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:bidi"), "ar-SA")
    lang.set(qn("w:val"), "ar-SA")


def set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, line=1.15):
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "right" if align == WD_ALIGN_PARAGRAPH.RIGHT else "center")


def add_p(container, text="", size=10.5, bold=False, color=CHARCOAL, align=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, keep=False):
    p = container.add_paragraph()
    set_paragraph_rtl(p, align, before, after)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_bullet(container, text, color=CHARCOAL, size=10.2, symbol="•", bold=False):
    p = container.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT, after=2)
    p.paragraph_format.right_indent = Cm(0.25)
    r1 = p.add_run(symbol + " ")
    set_run_font(r1, size, True, ORANGE)
    r2 = p.add_run(text)
    set_run_font(r2, size, bold, color)
    return p


def add_rule(doc, color=ORANGE, width=18):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(width))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_title(doc, number: str, title: str, subtitle: str | None = None):
    t = doc.add_table(rows=1, cols=2)
    set_table_rtl(t)
    t.autofit = False
    t.columns[0].width = Cm(2.0)
    t.columns[1].width = Cm(15.3)
    c_num, c_title = t.rows[0].cells
    set_cell_shading(c_num, ORANGE)
    set_cell_shading(c_title, LIGHT)
    for c in (c_num, c_title):
        set_cell_margins(c, 120, 160, 120, 160)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c_num.paragraphs[0]
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(number)
    set_run_font(r, 20, True, WHITE, rtl=False)
    p2 = c_title.paragraphs[0]
    set_paragraph_rtl(p2)
    r2 = p2.add_run(title)
    set_run_font(r2, 16, True, CHARCOAL)
    if subtitle:
        r3 = p2.add_run("\n" + subtitle)
        set_run_font(r3, 9.5, False, GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_info_box(doc, title, bullets: Sequence[str], fill=LIGHT_ORANGE, border=ORANGE, icon="✓"):
    t = doc.add_table(rows=1, cols=1)
    set_table_rtl(t)
    cell = t.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 160, 190, 160, 190)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "14")
        e.set(qn("w:color"), border)
        borders.append(e)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    set_paragraph_rtl(p)
    r = p.add_run(f"{icon}  {title}")
    set_run_font(r, 12, True, border)
    for b in bullets:
        add_bullet(cell, b, size=9.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, filename: str, width_cm=16.8, caption: str | None = None):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    if caption:
        cp = add_p(doc, caption, 8.5, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, after=6)
        cp.paragraph_format.keep_with_next = False


def add_two_column(doc, left_title, left_items, right_title, right_items, colors=(LIGHT_GREEN, LIGHT_RED)):
    t = doc.add_table(rows=1, cols=2)
    set_table_rtl(t)
    for idx, (title, items, fill, border) in enumerate([
        (right_title, right_items, colors[1], RED),
        (left_title, left_items, colors[0], GREEN),
    ]):
        c = t.cell(0, idx)
        set_cell_shading(c, fill)
        set_cell_margins(c, 130, 160, 130, 160)
        p = c.paragraphs[0]
        set_paragraph_rtl(p)
        r = p.add_run(title)
        set_run_font(r, 11.5, True, border)
        for it in items:
            add_bullet(c, it, size=9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_simple_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths=None, font_size=9.3):
    t = doc.add_table(rows=1, cols=len(headers))
    set_table_rtl(t)
    t.style = "Table Grid"
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_shading(c, ORANGE)
        set_cell_margins(c)
        p = c.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        set_run_font(r, font_size, True, WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            set_cell_shading(c, WHITE if ridx % 2 == 0 else LIGHT)
            set_cell_margins(c)
            p = c.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(str(value))
            set_run_font(r, font_size, False, CHARCOAL)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_header_footer(doc, guide_name: str):
    for section in doc.sections:
        header = section.header
        table = header.add_table(rows=1, cols=2, width=Cm(17))
        set_table_rtl(table)
        c1, c2 = table.rows[0].cells
        p = c1.paragraphs[0]
        set_paragraph_rtl(p)
        r = p.add_run("LiefeX")
        set_run_font(r, 12, True, ORANGE, rtl=False)
        p2 = c2.paragraphs[0]
        set_paragraph_rtl(p2)
        r2 = p2.add_run(guide_name)
        set_run_font(r2, 8.5, False, GRAY)
        footer = section.footer
        fp = footer.paragraphs[0]
        set_paragraph_rtl(fp, WD_ALIGN_PARAGRAPH.CENTER)
        rr = fp.add_run("LiefeX  |  النسخة المعتمدة  |  صفحة ")
        set_run_font(rr, 8.5, False, GRAY)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        fp._p.append(fld)


def configure_doc(doc: Document, title: str):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.55)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), FONT_NAME)
    doc.core_properties.title = title
    doc.core_properties.subject = "دليل تأهيل سائقي التوصيل لدى LiefeX"
    doc.core_properties.author = "LiefeX"


def cover(doc, guide_label: str, title: str, subtitle: str, image="flow.png"):
    t = doc.add_table(rows=1, cols=2)
    set_table_rtl(t)
    c1, c2 = t.rows[0].cells
    set_cell_shading(c1, WHITE)
    set_cell_shading(c2, WHITE)
    p = c1.paragraphs[0]
    set_paragraph_rtl(p)
    r = p.add_run("Liefe")
    set_run_font(r, 26, True, DARK, rtl=False)
    r2 = p.add_run("X")
    set_run_font(r2, 26, True, ORANGE, rtl=False)
    p2 = c2.paragraphs[0]
    set_paragraph_rtl(p2, WD_ALIGN_PARAGRAPH.CENTER)
    r3 = p2.add_run(guide_label)
    set_run_font(r3, 11, True, WHITE)
    set_cell_shading(c2, ORANGE)
    add_rule(doc, ORANGE, 20)
    add_p(doc, title, 25, True, CHARCOAL, WD_ALIGN_PARAGRAPH.RIGHT, before=18, after=4)
    add_p(doc, subtitle, 13, True, ORANGE_DARK, WD_ALIGN_PARAGRAPH.RIGHT, after=6)
    add_p(doc, "مسار سائقي التوصيل فقط", 10.5, False, GRAY, WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    add_image(doc, image, 17.0)
    add_info_box(doc, "هوية البرنامج", [
        "برنامج حضوري مهني لمدة 6 ساعات.",
        "تدريب نظري وتطبيقي، ثم اختبار نظري وعملي.",
        "إصدار بطاقة سائق معتمد بعد استكمال الأوراق والنجاح.",
    ], fill=CHARCOAL, border=ORANGE, icon="◆")
    add_p(doc, "LiefeX — الإصدار الإنتاجي النهائي", 8.5, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, before=8)
    add_page_break(doc)


def build_student():
    doc = Document()
    configure_doc(doc, "دليل تأهيل سائق التوصيل لدى LiefeX")
    cover(doc, "دليل المتدرب", "دورة تدريبية لتأهيل سائقي التوصيل لدى LiefeX", "الدليل الرسمي للمتدرب والمرجع التشغيلي للسائق المعتمد")
    add_header_footer(doc, "دليل المتدرب — تأهيل سائقي التوصيل")

    add_section_title(doc, "A", "هوية البرنامج والاعتماد", "ماذا يحصل عليه السائق وما المطلوب منه؟")
    add_image(doc, "accreditation.png", 17.0, "مخطط الاعتماد من إنشاء الحساب حتى إصدار البطاقة")
    add_simple_table(doc, ["البند", "المعتمد"], [
        ["اسم البرنامج", "دورة تدريبية لتأهيل سائقي التوصيل لدى LiefeX"],
        ["مدة الدورة", "6 ساعات شاملة الاستراحة والتقييم"],
        ["النجاح", "70% في الاختبار النظري و70% في الاختبار العملي"],
        ["الإعادة", "حتى 3 محاولات"],
        ["المخرج", "بطاقة سائق معتمد من LiefeX"],
        ["مكان التدريب", "يحدد بالتواصل مع السائق أو برسالة داخل التطبيق أو عبر الخط الساخن"],
    ], widths=[5.0, 12.0])
    add_info_box(doc, "هدف الدورة", [
        "استخدام التطبيق بصورة صحيحة ومنظمة.",
        "تنفيذ الطلب من القبول حتى التسليم بجودة ثابتة.",
        "حماية الطلب والتعامل المهني مع المتجر والعميل والإدارة.",
        "الاستجابة الآمنة للطوارئ والإساءة والمشكلات التشغيلية.",
    ], fill=LIGHT_GREEN, border=GREEN)
    add_page_break(doc)

    add_section_title(doc, "1", "شروط الاعتماد والاستعداد للدورة", "لا يُفعّل الحساب قبل اكتمال الدورة والمستندات")
    add_simple_table(doc, ["الشرط", "الحالة"], [[x, "إلزامي"] for x in [
        "حساب مكتمل على المنصة",
        "رخصة قيادة سارية",
        "هوية أو إقامة",
        "أوراق المركبة",
        "العمر 18 سنة أو أكثر",
        "وثيقة لا حكم عليه",
        "سند إقامة",
        "إتمام الدورة والنجاح",
    ]], widths=[12.2, 4.8])
    add_two_column(doc, "بعد النجاح", [
        "تراجع الإدارة المستندات والنتائج.",
        "يُفعّل الحساب بعد استكمال كل المتطلبات.",
        "تصدر بطاقة سائق معتمد من LiefeX.",
    ], "حماية الحساب والسرية", [
        "الحساب والهاتف ومعرّف الطلب وبيانات العميل معلومات حساسة.",
        "يمنع إعطاء الحساب أو كلمة المرور أو الهاتف لشخص آخر.",
        "تشغيل الحساب بواسطة شخص غير معتمد قد يؤدي إلى إغلاقه.",
    ])
    add_page_break(doc)

    add_section_title(doc, "2", "إنشاء الحساب وتسجيل الدخول", "التسجيل الصحيح هو أول خطوة في مسار الاعتماد")
    add_two_column(doc, "شاشة البداية", [
        "اختيار اللغة المناسبة.",
        "الدخول بالحساب المعتمد.",
        "زر «تسجيل كتاجر» يستخدم أيضاً لبدء تسجيل السائق بحسب الواجهة الحالية.",
    ], "شاشة التسجيل", [
        "اختيار نوع العمل: سائق توصيل.",
        "إدخال البيانات كاملة والموافقة على السياسات.",
        "انتظار مراجعة الإدارة بعد إنشاء الحساب.",
    ], colors=(LIGHT_GREEN, LIGHT_ORANGE))
    add_info_box(doc, "مهم", [
        "لا يمكن تفعيل «متاح» قبل موافقة الإدارة واستكمال المتطلبات.",
        "الاسم والعنوان لا يعدلان من داخل التطبيق؛ يراجع السائق الإدارة عند الحاجة.",
    ], fill=LIGHT_ORANGE, border=ORANGE, icon="!")
    add_page_break(doc)

    add_section_title(doc, "3", "لوحة السائق وبدء النوبة", "تفعيل «متاح» يعني الجاهزية الفعلية لاستقبال الطلبات")
    add_image(doc, "dashboard.png", 17.0, "شرح بصري للوحة السائق الرئيسية")
    add_two_column(doc, "قبل تشغيل «متاح»", [
        "الهاتف مشحون والإنترنت والموقع يعملان.",
        "الحقيبة والمركبة نظيفتان.",
        "معدات السلامة والفكة النقدية متوفرة.",
        "الزي وبطاقة التعريف مستخدمان.",
    ], "عند إنهاء النوبة", [
        "التحول إلى غير متاح عند التوقف المؤقت.",
        "تسجيل الخروج إلزامي عند انتهاء العمل.",
        "تسجيل الخروج يحول الحالة إلى Offline ويمنع الفوضى التشغيلية.",
    ])
    add_page_break(doc)

    add_section_title(doc, "4", "استقبال الطلب واتخاذ القرار", "المهلة المتاحة: دقيقة و20 ثانية")
    add_image(doc, "original.png", 17.0, "كرت طلب أصيل — تمثيل توضيحي مبني على الواجهة الفعلية")
    add_bullet(doc, "راجع اسم المتجر والكمية ومعرّف الطلب وطريقة الدفع قبل القبول.")
    add_bullet(doc, "القبول يغيّر حالة الطلب ويبلغ الأطراف ويبدأ مسار التنفيذ.")
    add_bullet(doc, "الرفض مسموح خلال المهلة، لكن التكرار غير المبرر يؤثر في التقييم والأولوية.")
    add_bullet(doc, "انتهاء المؤقت دون قرار يحول الطلب إلى «فائت» ويسجل في السجل.")
    add_info_box(doc, "قاعدة الجاهزية", [
        "من يفعّل «متاح» يعلن جاهزيته الفعلية للعمل.",
        "عند عدم الجاهزية: انتقل إلى غير متاح أو سجل الخروج؛ لا تتجاهل الطلبات.",
    ], fill=LIGHT_RED, border=RED, icon="!")
    add_page_break(doc)

    add_section_title(doc, "5", "الطلب الأصيل", "طلب منشأ داخل LiefeX ويحتوي بيانات المتجر والعميل")
    add_image(doc, "flow.png", 17.0)
    add_simple_table(doc, ["المرحلة", "ما يجب مراجعته"], [
        ["قبل القبول", "المتجر، الكمية، المعرّف، طريقة الدفع، الوقت"],
        ["بعد القبول", "عنوان المتجر، حالة الطلب، عنوان العميل، تفاصيل الطلب"],
        ["عند المتجر", "مطابقة معرّف الطلب وفحص التغليف خارجياً"],
        ["قبل العميل", "تأمين الطلب، تشغيل الملاحة، عدم الضغط على التسليم"],
        ["عند التسليم", "التحقق من المستلم والتحصيل النقدي عند الحاجة"],
    ], widths=[4.7, 12.3])
    add_page_break(doc)

    add_section_title(doc, "6", "الطلب الهجين", "طلب ينشئه المتجر لعميله الخاص غير المسجل في LiefeX")
    add_image(doc, "hybrid.png", 17.0, "في الطلب الهجين لا تظهر بيانات العميل كاملة في الكرت الأولي")
    add_info_box(doc, "لا تغادر المتجر قبل الحصول على البيانات الأربعة", [
        "اسم العميل.",
        "رقم هاتف العميل.",
        "العنوان الكامل والدقيق.",
        "قيمة المبلغ المطلوب تحصيله.",
    ], fill=LIGHT_RED, border=RED, icon="!")
    add_bullet(doc, "أدخل العنوان يدوياً في تطبيق الملاحة واختبره قبل الانطلاق.")
    add_bullet(doc, "أي معلومة ناقصة أو غير واضحة: توقف وتواصل مع الإدارة.")
    add_page_break(doc)

    add_section_title(doc, "7", "الاستلام المهني من المتجر", "التعريف، المطابقة، الفحص الخارجي، ثم التأمين")
    add_image(doc, "pickup.png", 17.0)
    add_info_box(doc, "الصيغة المهنية", [
        "أنا سائق ليفيكس ولدي طلب لديكم يتوجب علي استلامه.",
        "يبقى الهاتف مفتوحاً على الواجهة ويظهر السائق معرّف الطلب للموظف المختص فقط.",
    ], fill=LIGHT_GREEN, border=GREEN, icon="✓")
    add_two_column(doc, "افعل", [
        "استخدم سطحاً نظيفاً ومرتفعاً.",
        "افحص الختم والتسرب والتمزق خارجياً.",
        "أمّن الطلب في الحقيبة قبل تأكيد الاستلام.",
    ], "لا تفعل", [
        "لا تفتح الطلب المختوم.",
        "لا تضع الطلب على الأرض.",
        "لا تعطِ معرّف الطلب لشخص غير معني.",
    ])
    add_page_break(doc)

    add_section_title(doc, "8", "حماية الطلب والحقيبة", "السائق مسؤول عن سلامة الطلب أثناء النقل")
    add_image(doc, "bag.png", 17.0)
    add_bullet(doc, "حقيبة التوصيل إلزامية وتوفرها LiefeX للسائق خلال التدريب.")
    add_bullet(doc, "افصل الساخن عن البارد، وثبت المشروبات والسوائل.")
    add_bullet(doc, "استخدم وضعية مناسبة للبيتزا والكعك والمنتجات الهشة.")
    add_bullet(doc, "يمنع وضع الأغراض الشخصية مع الطلب.")
    add_bullet(doc, "نظف الحقيبة والمركبة دورياً وقبل العمل.")
    add_info_box(doc, "عند التلف أو الانسكاب", [
        "أبلغ الإدارة قبل التسليم.",
        "يعود السائق إلى المطعم الأصلي فور توجيه الإدارة.",
        "السائق مسؤول عن الضرر الناتج عن سوء النقل أو الإهمال.",
    ], fill=LIGHT_RED, border=RED, icon="!")
    add_page_break(doc)

    add_section_title(doc, "9", "تأكيد الاستلام والملاحة والتتبع", "لا تؤكد الاستلام إلا بعد تأمين الطلب والاستعداد للانطلاق")
    add_bullet(doc, "عند الضغط على رمز عنوان العميل تظهر رسالة: هل استلمت الطلب من المطعم وترغب بالتوجه للعميل؟")
    add_bullet(doc, "اختيار «نعم» يؤكد الاستلام، ويفتح الملاحة، ويبلغ العميل، ويفعل التتبع والوقت المتوقع.")
    add_bullet(doc, "التتبع لدى العميل لا يبدأ قبل تأكيد الاستلام والتوجه إليه.")
    add_info_box(doc, "قبل الضغط على نعم", [
        "استلم الطلب فعلياً.",
        "افحص التغليف خارجياً.",
        "ضع الطلب وثبته داخل الحقيبة.",
        "تأكد أنك جاهز للمغادرة فوراً.",
    ], fill=LIGHT_GREEN, border=GREEN)
    add_page_break(doc)

    add_section_title(doc, "10", "التعامل مع العميل والتسليم", "تواصل محدود، تحقق صحيح، وتسليم مهني")
    add_two_column(doc, "التواصل المسموح", [
        "الاتصال مرة واحدة فقط عند الحاجة للمساعدة في تحديد العنوان.",
        "العبارة المعتمدة: مرحباً، هنا سائق ليفيكس، هل يمكنك مساعدتي بالوصول لعنوان التسليم الصحيح؟ أنا الآن في موقع التسليم.",
        "انتظار العميل حتى 10 دقائق ثم التواصل مع الإدارة.",
    ], "ممنوع", [
        "أي تواصل خارج غرض تحديد عنوان التسليم.",
        "طلب التقييم أو الإكرامية.",
        "الجدال أو الوعد بتعويض أو إعادة الطلب دون تعليمات.",
    ])
    add_bullet(doc, "تحقق من المستلم بواسطة معرّف الطلب.")
    add_bullet(doc, "يمكن التسليم لشخص آخر بعد تأكيد العميل.")
    add_bullet(doc, "الترك عند الباب لا يتم إلا بتأكيد العميل وتوثيق صورة.")
    add_bullet(doc, "عند رفض العميل: اتصل بالإدارة وأعد الطلب فوراً إلى الحقيبة.")
    add_page_break(doc)

    add_section_title(doc, "11", "الدفع النقدي والتسوية", "حصّل قيمة الطلب فقط، ثم سلّم الكاش كاملاً للإدارة")
    add_image(doc, "cash.png", 17.0)
    add_bullet(doc, "طريقة الدفع تظهر على كرت الطلب.")
    add_bullet(doc, "في الطلب النقدي يحصّل السائق قيمة الطلب الظاهرة فقط، وليس أجرة التوصيل.")
    add_bullet(doc, "يجب توفر فكة كافية قبل تفعيل «متاح».")
    add_bullet(doc, "في حال الرفض أو النقص في الدفع: حافظ على الهدوء واتصل بالإدارة فوراً.")
    add_bullet(doc, "في نهاية كل نوبة يسلم السائق كامل الكاش المستحق إلى الإدارة ولا يخصم مستحقاته بنفسه.")
    add_bullet(doc, "عمولة LiefeX البالغة 5% تخصم من أرباح السائق وتظهر في كشف التسوية لدى الإدارة.")
    add_page_break(doc)

    add_section_title(doc, "12", "زر التسليم وإنهاء دورة الطلب", "الضغط المبكر مخالفة خطرة لأنه يغلق الطلب ويوقف التتبع")
    add_info_box(doc, "اضغط «تم التسليم» فقط بعد", [
        "تسليم الطلب فعلياً إلى الشخص الصحيح.",
        "التحقق من معرّف الطلب أو تأكيد العميل.",
        "تحصيل قيمة الطلب عند الدفع النقدي.",
    ], fill=LIGHT_GREEN, border=GREEN)
    add_info_box(doc, "ماذا يحدث بعد الضغط؟", [
        "تتحول الحالة إلى تم التسليم.",
        "تُبلغ جميع الأطراف.",
        "تحسب المسافات وتكلفة التوصيل.",
        "يسجل استحقاق السائق في الأرباح والسجل.",
    ], fill=LIGHT_ORANGE, border=ORANGE, icon="→")
    add_page_break(doc)

    add_section_title(doc, "13", "الأرباح والضبط والسجل", "استخدم البيانات للتوثيق والتسوية ومراجعة الأداء")
    add_simple_table(doc, ["التبويب", "الوظيفة"], [
        ["الأرباح", "عدد التوصيلات، الأرباح اليومية، المقبول والملغي والمنتهي، المسافات والإجمالي"],
        ["الضبط", "الملف الشخصي، نطاق العمل، التسعيرة المعتمدة، تفعيل الدفع"],
        ["السجل", "أرشيف الطلبات، الحالات، المسافات، التاريخ، الفلاتر، الطباعة والتحميل"],
    ], widths=[4.0, 13.0])
    add_info_box(doc, "تمييز مالي مهم", [
        "الأرباح: حق السائق مقابل التوصيلات.",
        "الكاش المستحق: أموال طلبات نقدية يجب تسليمها للإدارة.",
        "خصم 5% يظهر في كشف التسوية، وليس كبند داخل تبويب الأرباح.",
    ], fill=LIGHT_ORANGE, border=ORANGE, icon="$")
    add_bullet(doc, "الحد الأدنى لنطاق العمل 3 كم، والسائق يختار نطاقه ضمن المتاح.")
    add_bullet(doc, "التقييم بالنجوم صادر عن العملاء والمتاجر.")
    add_page_break(doc)

    add_section_title(doc, "14", "السلامة والمركبة والزي", "السلامة جزء من جودة الخدمة وشروط العمل")
    add_image(doc, "safety.png", 17.0)
    add_bullet(doc, "الزي الرسمي وبطاقة التعريف إلزاميان أثناء العمل.")
    add_bullet(doc, "الخوذة إلزامية لسائق الدراجة النارية أثناء الرحلة.")
    add_bullet(doc, "يمنع استخدام الهاتف أثناء القيادة.")
    add_bullet(doc, "يمنع اصطحاب أشخاص آخرين أثناء العمل.")
    add_bullet(doc, "يجب أن تتوفر معدات سلامة المركبة وأن تكون المركبة نظيفة وصالحة.")
    add_bullet(doc, "المركبات المقبولة: سيارة، دراجة نارية، دراجة هوائية.")
    add_page_break(doc)

    add_section_title(doc, "15", "الحالات الطارئة والمشكلات", "السلامة أولاً، ثم التواصل مع الإدارة والجهات المختصة")
    add_image(doc, "emergency.png", 17.0)
    add_simple_table(doc, ["الحالة", "الإجراء"], [
        ["حادث", "تأمين الموقع، الاتصال بالإدارة والجهات المختصة"],
        ["عطل المركبة", "الاتصال بالإدارة لتأمين بديل ومعالجة الطلب"],
        ["مرض أو إصابة", "التوقف في مكان آمن وتأمين النفس والمركبة والاتصال بالإدارة"],
        ["تعطل التطبيق", "اتصال هاتفي مباشر بالإدارة"],
        ["ضياع أو سرقة الهاتف", "استخدام أي جهاز متاح للاتصال بالإدارة وإبلاغ الموقع"],
        ["طريق مغلق", "البحث عن بدائل، وعند الاستحالة التواصل مع الإدارة"],
        ["خطر أمني", "الانتقال إلى مكان آمن والاتصال بالإدارة والجهات المختصة"],
    ], widths=[4.3, 12.7], font_size=8.8)
    add_page_break(doc)

    add_section_title(doc, "16", "الإساءة والبلاغات والأدلة", "لا ترد بالمثل؛ احمِ نفسك ووثق الواقعة")
    add_bullet(doc, "تشمل الإساءة: السب، التهديد، التحرش، التمييز، الاعتداء، الاحتيال أو إجبار السائق على مخالفة.")
    add_bullet(doc, "إن كان إكمال الطلب آمناً، أكمله ثم أبلغ الإدارة.")
    add_bullet(doc, "عند وجود خطر شديد: أوقف العمل وانتقل إلى مكان آمن واتصل بالإدارة والجهات المختصة.")
    add_bullet(doc, "استخدم الدعم داخل التطبيق، والاتصال المباشر عند تعطل التطبيق أو الخطر العاجل.")
    add_bullet(doc, "وثق بصورة أو فيديو متى كان ذلك آمناً.")
    add_bullet(doc, "تتحقق الإدارة من البلاغات بعدل ولا تعاقب السائق تلقائياً بناء على شكوى غير مثبتة.")
    add_info_box(doc, "نزاهة البلاغ", ["تقديم صورة أو فيديو أو بلاغ مزور عمداً يؤدي إلى الرسوب المباشر وإجراء إداري."], fill=LIGHT_RED, border=RED, icon="!")
    add_page_break(doc)

    add_section_title(doc, "17", "حالات الطلب وألوانها", "اقرأ الإطار والحالة قبل اتخاذ أي إجراء")
    add_simple_table(doc, ["الحالة", "لون الإطار", "المعنى"], [
        ["على الطريق", "اللون التشغيلي في التطبيق", "الطلب قيد التنفيذ"],
        ["تم التسليم", "أخضر", "اكتملت دورة الطلب"],
        ["فائت", "رمادي", "انتهت المهلة دون قرار"],
        ["ملغي", "أحمر", "ألغي الطلب قبل القبول أو بإجراء إداري"],
    ], widths=[4.2, 4.2, 8.6])
    add_info_box(doc, "بعد القبول", [
        "لا يستطيع السائق إلغاء الطلب مستقلاً.",
        "في الطوارئ يتواصل مع الإدارة أو يكمل دورة الطلب وفق التوجيه.",
        "الإدارة وحدها تتولى التدخلات الضرورية وإعادة التعيين.",
    ], fill=LIGHT_ORANGE, border=ORANGE, icon="!")
    add_page_break(doc)

    add_section_title(doc, "18", "المخالفات التي تؤدي إلى الرسوب المباشر", "الخطأ الحرج يلغي نتيجة المحاولة العملية الحالية")
    direct_fail = [
        "مشاركة الحساب أو الهاتف أو السماح لشخص آخر باستخدام التطبيق.",
        "كشف بيانات العميل أو معرّف الطلب لغير المعنيين.",
        "فتح الطلب المختوم أو وضعه على الأرض.",
        "مغادرة المتجر قبل تأمين الطلب داخل الحقيبة.",
        "مغادرة طلب هجين دون بيانات العميل الأربعة أو دون اختبار العنوان.",
        "الضغط على تم التسليم قبل التسليم الفعلي أو قبل التحصيل النقدي.",
        "تسليم الطلب إلى شخص غير صحيح دون تحقق أو تأكيد.",
        "إخفاء تلف أو فقدان الطلب وعدم إبلاغ الإدارة.",
        "استخدام الهاتف أثناء القيادة أو القيادة دون خوذة للدراجة النارية.",
        "العمل بمركبة غير آمنة أو اصطحاب شخص آخر.",
        "الرد على الإساءة باعتداء أو تقديم دليل مزور.",
        "أي تصرف يعرض السائق أو الآخرين لخطر مباشر.",
    ]
    for item in direct_fail:
        add_bullet(doc, item, color=RED, symbol="✕")
    add_info_box(doc, "ملاحظة", ["الرسوب المباشر يلغي المحاولة الحالية، ولا يعني إغلاق الحساب تلقائياً. يسمح بالإعادة حتى 3 مرات."], fill=LIGHT_ORANGE, border=ORANGE, icon="i")
    add_page_break(doc)

    add_section_title(doc, "19", "قائمة مراجعة السائق", "استخدمها قبل النوبة وأثناء الطلب وعند إنهاء العمل")
    add_simple_table(doc, ["التوقيت", "قائمة التحقق"], [
        ["قبل النوبة", "حساب معتمد — هاتف مشحون — موقع وإنترنت — زي وبطاقة — حقيبة نظيفة — سلامة المركبة — فكة"],
        ["عند الطلب", "مراجعة البيانات — قرار خلال 1:20 — عدم التجاهل"],
        ["عند المتجر", "تعريف — معرّف — فحص خارجي — سطح مرتفع — تأمين الحقيبة"],
        ["الطلب الهجين", "اسم — هاتف — عنوان كامل — مبلغ التحصيل — اختبار الملاحة"],
        ["عند العميل", "تحقق — تحصيل — تسليم فعلي — ثم تأكيد التسليم"],
        ["نهاية النوبة", "غير متاح — تسجيل خروج — تسليم الكاش — مراجعة التسوية"],
    ], widths=[3.5, 13.5], font_size=8.8)
    add_info_box(doc, "قاعدة LiefeX الذهبية", ["لا تؤكد أي مرحلة في التطبيق قبل تنفيذها فعلياً على أرض الواقع."], fill=CHARCOAL, border=ORANGE, icon="◆")
    add_page_break(doc)

    add_section_title(doc, "20", "إقرار المتدرب", "توقيع المتدرب يثبت استلام المحتوى وفهم القواعد")
    add_p(doc, "أقر أنا المتدرب بأنني حضرت دورة تأهيل سائقي التوصيل لدى LiefeX، واطلعت على قواعد استخدام التطبيق، واستلام الطلب وحمايته، والتعامل مع المتجر والعميل، والدفع النقدي، والسلامة والطوارئ، وأفهم أن الاعتماد قابل للإلغاء عند ارتكاب مخالفات تستوجب ذلك.", 11, False, CHARCOAL, after=16)
    for label in ["اسم المتدرب:", "رقم الحساب:", "التاريخ:", "التوقيع:", "اسم المدرب:", "توقيع الإدارة:"]:
        add_p(doc, label + "  ........................................................................................................", 10.5, True, CHARCOAL, after=10)
    add_info_box(doc, "إصدار البطاقة", ["تصدر بطاقة سائق معتمد من LiefeX بعد استكمال الأوراق والنجاح في الاختبار النظري والعملي."], fill=LIGHT_GREEN, border=GREEN)

    path = OUT / "دليل_تأهيل_سائق_التوصيل_LiefeX_النسخة_الإنتاجية.docx"
    doc.save(path)
    embed_fonts(path)
    return path


def q(text, options, correct, points):
    return {"text": text, "options": options, "correct": correct, "points": points}


def build_trainer():
    doc = Document()
    configure_doc(doc, "دليل المدرب والاختبارات لدى LiefeX")
    cover(doc, "دليل المدرب", "دليل المدرب والاختبارات", "خطة تقديم دورة تأهيل سائقي التوصيل وتقييم الاعتماد", image="accreditation.png")
    add_header_footer(doc, "دليل المدرب والاختبارات — LiefeX")

    add_section_title(doc, "A", "طريقة استخدام الدليل", "مرجع موحد للتدريب والتقييم واتخاذ قرار الاعتماد")
    add_two_column(doc, "مسؤوليات المدرب", [
        "تقديم القواعد كما هي دون اختصار مخل أو إضافة غير معتمدة.",
        "إظهار التطبيق ومحاكاة الطلبات عملياً.",
        "التأكد من فهم المتدرب قبل الانتقال بين الوحدات.",
        "توثيق الدرجات والأخطاء الحرجة بدقة.",
    ], "ممنوع أثناء التدريب", [
        "إعطاء المتدرب الإجابات قبل الاختبار.",
        "تجاوز الاختبار العملي أو اعتباره اختيارياً.",
        "تمرير متدرب ارتكب خطأ رسوب مباشر.",
        "تفعيل الحساب قبل اكتمال الأوراق والنتائج.",
    ])
    add_info_box(doc, "نتيجة البرنامج", [
        "ناجح: 70% فأكثر في النظري و70% فأكثر في العملي دون خطأ مباشر.",
        "غير ناجح: يعاد الاختبار ضمن حد أقصى 3 محاولات.",
        "التفعيل: بعد الدورة واستكمال الأوراق واعتماد الإدارة.",
    ], fill=LIGHT_GREEN, border=GREEN)
    add_page_break(doc)

    add_section_title(doc, "1", "الجدول الزمني للدورة — 6 ساعات", "وقت كل وحدة محسوب ضمن اليوم التدريبي")
    add_simple_table(doc, ["الفترة", "المدة", "المحتوى"], [
        ["الوحدة الأولى", "1 ساعة و30 دقيقة", "الحساب، الاعتماد، التطبيق، التوفر، استقبال الطلبات"],
        ["استراحة", "30 دقيقة", "استراحة منظمة"],
        ["الوحدة الثانية", "1 ساعة و30 دقيقة", "الأصيل والهجين، المتجر، حماية الطلب، العميل"],
        ["شرح ومحاكاة", "45 دقيقة", "تطبيق عملي على الطلبات والدفع والطوارئ"],
        ["الاختبار العملي", "45 دقيقة", "أصيل، هجين، نقدي، طارئ"],
        ["الاستفسارات", "15 دقيقة", "مراجعة أخيرة قبل النظري"],
        ["الاختبار النظري", "45 دقيقة", "15 سؤال اختيار من متعدد"],
    ], widths=[3.2, 3.8, 10.0], font_size=8.7)
    add_image(doc, "flow.png", 17.0)
    add_page_break(doc)

    add_section_title(doc, "2", "خطة الوحدة الأولى", "الحساب والتطبيق والتوفر واستقبال الطلب")
    add_simple_table(doc, ["الموضوع", "أسلوب العرض", "دليل التحقق"], [
        ["الاعتماد", "اشرح الشروط ومسار التفعيل والبطاقة", "يعدد المتدرب المتطلبات دون مساعدة"],
        ["التسجيل", "عرض شاشة البداية والتسجيل", "يختار سائق توصيل ويشرح انتظار الموافقة"],
        ["لوحة السائق", "عرض الدعم والتوفر والتبويبات والخروج", "يميز بين متاح وغير متاح"],
        ["كرت الطلب", "عرض المؤقت والقبول والرفض والفائت", "يتخذ قراراً صحيحاً خلال المهلة"],
        ["الحالات", "شرح الأخضر والرمادي والأحمر", "يربط اللون بالحالة الصحيحة"],
    ], widths=[3.5, 7.4, 6.1], font_size=8.6)
    add_image(doc, "dashboard.png", 16.5)
    add_page_break(doc)

    add_section_title(doc, "3", "خطة الوحدة الثانية", "تنفيذ الطلب وحمايته والتعامل مع الأطراف")
    add_simple_table(doc, ["الموضوع", "النقطة الأساسية", "تمرين سريع"], [
        ["الطلب الأصيل", "مراجعة البيانات في كل مرحلة", "يشرح المسار من القبول حتى التسليم"],
        ["الطلب الهجين", "البيانات الأربعة واختبار الملاحة", "يستخرج النواقص من سيناريو"],
        ["المتجر", "تعريف ومطابقة وفحص خارجي وتأمين", "محاكاة الاستلام على طاولة"],
        ["الحقيبة", "فصل الساخن والبارد وتثبيت السوائل", "ترتيب طلب نموذجي"],
        ["العميل", "تواصل محدود وتحقق وانتظار 10 دقائق", "مكالمة إرشاد للعنوان"],
        ["النقدي", "تحصيل كامل وتسليم الكاش للإدارة", "احتساب فكة وتسوية"],
    ], widths=[3.4, 7.2, 6.4], font_size=8.5)
    add_image(doc, "pickup.png", 16.5)
    add_page_break(doc)

    add_section_title(doc, "4", "إعداد بيئة التدريب", "المواد والأدوات المطلوبة قبل حضور المتدربين")
    add_simple_table(doc, ["الأداة", "الغرض"], [
        ["هاتف تجريبي أو حساب تدريب", "عرض التطبيق وتنفيذ المحاكاة"],
        ["حقيبة LiefeX", "تدريب التعبئة والفصل والتثبيت"],
        ["عبوات ساخنة وباردة ومشروبات تجريبية", "تقييم حماية الطلب"],
        ["مبلغ نقدي وفكة تدريبية", "محاكاة الطلب النقدي"],
        ["طاولة نظيفة", "محاكاة الاستلام من المتجر"],
        ["نماذج الاختبار والتقييم", "توثيق النتائج"],
        ["بطاقة تعريف وزي وخوذة", "شرح المظهر والسلامة"],
    ], widths=[7.0, 10.0])
    add_info_box(doc, "قبل بدء الدورة", [
        "تأكد من شحن الأجهزة وعمل الإنترنت والموقع.",
        "حضّر الطلبات التجريبية والعناوين والسيناريوهات.",
        "لا تستخدم بيانات عميل حقيقي في التدريب.",
    ], fill=LIGHT_ORANGE, border=ORANGE)
    add_page_break(doc)

    add_section_title(doc, "5", "سيناريو الاختبار العملي — الطلب الأصيل", "درجة المحور: 20 نقطة")
    add_simple_table(doc, ["المعيار", "النقاط", "النجاح"], [
        ["مراجعة بيانات الطلب واتخاذ القرار", "4", "قرار صحيح ضمن المهلة"],
        ["الوصول للمتجر وإظهار المعرّف", "4", "تعريف مهني وسرية"],
        ["الفحص الخارجي وتأمين الطلب", "5", "لا فتح ولا أرض، وتثبيت صحيح"],
        ["تأكيد الاستلام والملاحة", "3", "بعد الجاهزية الفعلية"],
        ["التسليم والتحقق والإغلاق", "4", "تسليم فعلي ثم تأكيد"],
    ], widths=[8.0, 2.5, 6.5])
    add_image(doc, "original.png", 16.5)
    add_page_break(doc)

    add_section_title(doc, "6", "سيناريو الاختبار العملي — الطلب الهجين", "درجة المحور: 20 نقطة")
    add_simple_table(doc, ["المعيار", "النقاط", "النجاح"], [
        ["يميز الطلب الهجين", "3", "يشرح اختلافه عن الأصيل"],
        ["يحصل على اسم العميل ورقمه", "4", "دون نقص"],
        ["يحصل على العنوان الكامل", "4", "عنوان واضح ودقيق"],
        ["يحصل على قيمة التحصيل", "3", "يكررها للتأكيد"],
        ["يدخل العنوان ويختبر الملاحة", "4", "قبل مغادرة المتجر"],
        ["يتوقف عند نقص المعلومات", "2", "يتواصل مع الإدارة"],
    ], widths=[8.0, 2.5, 6.5])
    add_image(doc, "hybrid.png", 16.5)
    add_page_break(doc)

    add_section_title(doc, "7", "سيناريو الاختبار العملي — الطلب النقدي", "درجة المحور: 15 نقطة")
    add_image(doc, "cash.png", 16.5)
    add_simple_table(doc, ["المعيار", "النقاط"], [
        ["يحدد قيمة الطلب المطلوبة من الكرت", "3"],
        ["يميز قيمة الطلب عن أجرة التوصيل وأرباحه", "3"],
        ["يمتلك فكة ويحصّل المبلغ كاملاً", "4"],
        ["لا يؤكد التسليم قبل التحصيل", "3"],
        ["يسلم الكاش كاملاً للإدارة ولا يخصم ذاتياً", "2"],
    ], widths=[13.5, 3.5])
    add_page_break(doc)

    add_section_title(doc, "8", "سيناريو الاختبار العملي — حالة طارئة", "درجة المحور: 15 نقطة")
    add_image(doc, "emergency.png", 16.5)
    add_simple_table(doc, ["المعيار", "النقاط"], [
        ["يؤمّن نفسه والمركبة والموقع", "4"],
        ["يحافظ على الطلب داخل الحقيبة", "2"],
        ["يتصل بالإدارة ويعطي معلومات واضحة", "4"],
        ["يتصل بالجهات المختصة عند الحاجة", "3"],
        ["لا يتصرف منفرداً في إعادة التعيين أو التعويض", "2"],
    ], widths=[13.5, 3.5])
    add_page_break(doc)

    add_section_title(doc, "9", "نموذج التقييم العملي — 100 نقطة", "النجاح من 70 مع عدم وجود خطأ رسوب مباشر")
    practical = [
        ["الاستعداد والحساب والملف", "10", ""],
        ["التوفر واستقبال الطلب واتخاذ القرار", "10", ""],
        ["تنفيذ طلب أصيل", "20", ""],
        ["تنفيذ طلب هجين", "20", ""],
        ["الاستلام وحماية الطلب", "10", ""],
        ["طلب نقدي وتسوية", "15", ""],
        ["حالة طارئة أو إساءة", "15", ""],
        ["المجموع", "100", ""],
    ]
    add_simple_table(doc, ["المحور", "الدرجة القصوى", "درجة المتدرب"], practical, widths=[10.0, 3.5, 3.5])
    add_p(doc, "اسم المتدرب: ............................................................................................", 10.5, True, after=8)
    add_p(doc, "اسم المقيم: ...............................................................................................", 10.5, True, after=8)
    add_p(doc, "النتيجة:  ناجح  ☐     إعادة  ☐     رسوب مباشر  ☐", 10.5, True, after=8)
    add_p(doc, "الملاحظات: ..............................................................................................................................................................", 10.5, True, after=8)
    add_page_break(doc)

    questions = [
        q("متى يجوز للسائق تفعيل حالة متاح؟", ["بعد فتح التطبيق مباشرة", "عند الجاهزية الفعلية واستكمال الاعتماد", "عند وصول أول طلب", "في نهاية النوبة"], 1, 6),
        q("ماذا يفعل السائق عند انتهاء النوبة؟", ["يترك التطبيق مفتوحاً", "يحذف السجل", "يسجل الخروج", "يرفض الطلب التالي"], 2, 7),
        q("ماذا يحدث عند انتهاء مؤقت الطلب دون قرار؟", ["يقبل تلقائياً", "يتحول إلى فائت", "يلغى الحساب", "يختفي دون تسجيل"], 1, 7),
        q("ما الذي يراجعه السائق قبل قبول الطلب الأصيل؟", ["اسم المتجر والكمية والمعرّف وطريقة الدفع", "التقييم فقط", "اسم السائق", "كشف التسوية"], 0, 7),
        q("ما البيانات الإلزامية قبل مغادرة المتجر في الطلب الهجين؟", ["اسم العميل فقط", "العنوان فقط", "الاسم والهاتف والعنوان الكامل وقيمة التحصيل", "رقم المركبة"], 2, 7),
        q("لمن يجوز إظهار معرّف الطلب؟", ["أي شخص في المتجر", "الموظف المختص باستلام الطلب فقط", "أصدقاء السائق", "عميل آخر"], 1, 7),
        q("ما نوع فحص الطلب المسموح للسائق؟", ["فتح العبوة وعد المحتويات", "تذوق المنتج", "فحص خارجي للختم والتسرب والتمزق", "إزالة الختم"], 2, 7),
        q("كيف تحمى الطلبات الساخنة والباردة؟", ["توضع معاً", "تفصل وتثبت داخل الحقيبة", "توضع على المقعد", "تترك دون حقيبة"], 1, 6),
        q("في الطلب النقدي، ماذا يحصّل السائق؟", ["قيمة الطلب الظاهرة فقط", "أجرة التوصيل فقط", "ربحه فقط", "أي مبلغ يختاره"], 0, 7),
        q("متى يضغط السائق نعم لتأكيد الاستلام من المتجر؟", ["قبل وصوله", "بعد تأمين الطلب والاستعداد للانطلاق", "بعد التسليم", "عند فتح التطبيق"], 1, 6),
        q("متى يجوز الضغط على تم التسليم؟", ["بعد القبول", "عند الوصول للشارع", "بعد التسليم الفعلي والتحصيل عند الحاجة", "قبل الاتصال بالعميل"], 2, 7),
        q("ما الإجراء الأول عند حادث أو خطر؟", ["إغلاق الحساب", "تأمين النفس والموقع ثم الاتصال بالإدارة والجهات المختصة", "ترك الطلب", "الجدال مع الطرف الآخر"], 1, 7),
        q("كيف يتصرف السائق عند الإساءة؟", ["يرد بالمثل", "يوثق ويتواصل مع الإدارة ويحمي نفسه", "ينشر بيانات العميل", "يلغي الحساب"], 1, 6),
        q("أين تظهر عمولة LiefeX البالغة 5%؟", ["داخل كرت الطلب", "في كشف التسوية لدى الإدارة", "في عنوان العميل", "لا تظهر أبداً"], 1, 6),
        q("أي عبارة صحيحة أثناء القيادة؟", ["يجوز استخدام الهاتف باليد", "يجوز اصطحاب مرافق", "الخوذة ومعدات السلامة إلزامية والهاتف ممنوع", "السلامة اختيارية"], 2, 7),
    ]

    add_section_title(doc, "10", "تعليمات الاختبار النظري", "15 سؤالاً — 45 دقيقة — المجموع 100")
    add_bullet(doc, "جميع الأسئلة من نوع اختر الإجابة الصحيحة.")
    add_bullet(doc, "لا يسمح باستخدام الهاتف أو الدليل أثناء الاختبار.")
    add_bullet(doc, "تكتب الإجابة في ورقة مستقلة أو توضع علامة واضحة أمام الخيار.")
    add_bullet(doc, "النجاح من 70 نقطة.")
    add_p(doc, "اسم المتدرب: ...............................................................................   التاريخ: ...........................", 10.5, True, after=8)
    add_page_break(doc)

    # questions across pages
    for idx, item in enumerate(questions, 1):
        if idx in (1, 6, 11):
            add_section_title(doc, f"Q{(idx-1)//5+1}", f"أسئلة الاختبار النظري — {idx} إلى {min(idx+4,15)}", "اختر إجابة واحدة فقط")
        p = add_p(doc, f"{idx}. {item['text']}  ({item['points']} درجات)", 10.7, True, CHARCOAL, after=3)
        for letter, opt in zip(["أ", "ب", "ج", "د"], item["options"]):
            add_p(doc, f"☐ {letter}) {opt}", 9.8, False, CHARCOAL, after=1)
        add_p(doc, "", 4, after=3)
        if idx in (5, 10):
            add_page_break(doc)
    add_page_break(doc)

    add_section_title(doc, "11", "مفتاح الإجابة وتوزيع الدرجات", "يستخدمه المدرب بعد استلام أوراق المتدربين")
    answer_rows = []
    letters = ["أ", "ب", "ج", "د"]
    for i, item in enumerate(questions, 1):
        answer_rows.append([str(i), letters[item["correct"]], str(item["points"])])
    add_simple_table(doc, ["السؤال", "الإجابة", "الدرجة"], answer_rows, widths=[5.5, 5.5, 5.5])
    add_p(doc, "المجموع: .......... / 100     النتيجة: ناجح ☐   إعادة ☐", 11, True, after=8)
    add_page_break(doc)

    add_section_title(doc, "12", "أخطاء الرسوب المباشر", "وجود أي خطأ منها يلغي المحاولة العملية الحالية")
    for item in [
        "مشاركة الحساب أو الهاتف أو كلمة المرور.",
        "فتح الطلب المختوم أو وضعه على الأرض.",
        "مغادرة المتجر دون تأمين الطلب.",
        "مغادرة الطلب الهجين دون البيانات الأربعة أو دون اختبار العنوان.",
        "الضغط على تم التسليم قبل التسليم الفعلي أو قبل التحصيل النقدي.",
        "تسليم الطلب لشخص غير صحيح دون تحقق.",
        "إخفاء التلف أو الفقد وعدم إبلاغ الإدارة.",
        "استخدام الهاتف أثناء القيادة أو عدم ارتداء الخوذة للدراجة النارية.",
        "مركبة غير آمنة أو اصطحاب مرافق أثناء العمل.",
        "الرد على الإساءة باعتداء أو تزوير دليل.",
        "أي تصرف يعرض الأشخاص لخطر مباشر.",
    ]:
        add_bullet(doc, item, color=RED, symbol="✕")
    add_info_box(doc, "قرار المدرب", ["يسجل الخطأ بالوصف والوقت والسيناريو، ويبلغ المتدرب أن النتيجة تخص المحاولة الحالية. يسمح بالإعادة حتى 3 مرات."], fill=LIGHT_RED, border=RED)
    add_page_break(doc)

    add_section_title(doc, "13", "سجل المحاولات وإعادة الاختبار", "الحد الأقصى: 3 محاولات")
    add_simple_table(doc, ["المحاولة", "التاريخ", "النظري", "العملي", "خطأ مباشر؟", "النتيجة"], [
        ["الأولى", "", "", "", "", ""],
        ["الثانية", "", "", "", "", ""],
        ["الثالثة", "", "", "", "", ""],
    ], widths=[2.7, 3.2, 2.4, 2.4, 3.2, 3.1], font_size=8.5)
    add_p(doc, "سبب الإعادة أو الرسوب: ........................................................................................................................................................", 10, True, after=12)
    add_p(doc, "خطة التصحيح قبل المحاولة التالية: ........................................................................................................................................", 10, True, after=12)
    add_page_break(doc)

    add_section_title(doc, "14", "قائمة مراجعة المستندات والتفعيل", "تراجعها الإدارة بعد الدورة وقبل تفعيل الحساب")
    docs = ["حساب مكتمل", "رخصة قيادة", "هوية أو إقامة", "أوراق المركبة", "العمر 18+", "لا حكم عليه", "سند إقامة", "نجاح نظري", "نجاح عملي"]
    add_simple_table(doc, ["المتطلب", "مكتمل", "ملاحظة"], [[d, "☐", ""] for d in docs], widths=[8.0, 3.0, 6.0])
    add_p(doc, "قرار الإدارة: تفعيل الحساب ☐    استكمال نواقص ☐    عدم اعتماد ☐", 10.5, True, after=10)
    add_p(doc, "اسم المراجع: .................................................... التوقيع: .................................................... التاريخ: ............................", 10, True)
    add_page_break(doc)

    add_section_title(doc, "15", "نموذج إصدار بطاقة السائق المعتمد", "لا تصدر البطاقة قبل قرار الإدارة النهائي")
    for label in [
        "اسم السائق:", "رقم الحساب:", "رقم بطاقة الاعتماد:", "نوع المركبة:", "تاريخ الاعتماد:", "اسم مسؤول الإصدار:", "توقيع السائق عند الاستلام:",
    ]:
        add_p(doc, label + "  ............................................................................................................................", 10.5, True, after=9)
    add_info_box(doc, "صلاحية الاعتماد", [
        "الاعتماد دائم ما لم ترتكب مخالفات تقضي الإدارة بسببها بإلغاء الاعتماد أو إغلاق الحساب.",
        "تخضع المركبة والحقيبة والنظافة والسلامة للتفتيش الدوري.",
    ], fill=LIGHT_GREEN, border=GREEN)

    path = OUT / "دليل_المدرب_والاختبارات_LiefeX_النسخة_الإنتاجية.docx"
    doc.save(path)
    embed_fonts(path)
    return path


def embed_fonts(docx_path: Path):
    """Embed Almarai regular and bold in DOCX using OOXML font embedding."""
    if not REGULAR_FONT.exists() or not BOLD_FONT.exists():
        return
    temp = docx_path.with_suffix(".tmp.docx")
    guid = uuid.uuid4()
    key = guid.bytes_le

    def obfuscate(data: bytes) -> bytes:
        b = bytearray(data)
        for i in range(min(32, len(b))):
            b[i] ^= key[i % 16]
        return bytes(b)

    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as zout:
        names = set(zin.namelist())
        for item in zin.infolist():
            if item.filename in {"word/fontTable.xml", "word/_rels/fontTable.xml.rels", "[Content_Types].xml"}:
                continue
            zout.writestr(item, zin.read(item.filename))

        from lxml import etree
        ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        ns_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
        font_xml = zin.read("word/fontTable.xml") if "word/fontTable.xml" in names else b'<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
        root = etree.fromstring(font_xml)
        # remove existing Almarai entry
        for f in root.xpath('.//w:font[@w:name="Almarai"]', namespaces={'w': ns_w}):
            f.getparent().remove(f)
        font_el = etree.SubElement(root, f"{{{ns_w}}}font")
        font_el.set(f"{{{ns_w}}}name", FONT_NAME)
        etree.SubElement(font_el, f"{{{ns_w}}}family").set(f"{{{ns_w}}}val", "swiss")
        etree.SubElement(font_el, f"{{{ns_w}}}charset").set(f"{{{ns_w}}}val", "B2")
        er = etree.SubElement(font_el, f"{{{ns_w}}}embedRegular")
        er.set(f"{{{ns_r}}}id", "rIdAlmaraiRegular")
        er.set(f"{{{ns_w}}}fontKey", "{" + str(guid).upper() + "}")
        eb = etree.SubElement(font_el, f"{{{ns_w}}}embedBold")
        eb.set(f"{{{ns_r}}}id", "rIdAlmaraiBold")
        eb.set(f"{{{ns_w}}}fontKey", "{" + str(guid).upper() + "}")
        zout.writestr("word/fontTable.xml", etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes"))

        rel_root = etree.Element(f"{{{ns_rel}}}Relationships")
        if "word/_rels/fontTable.xml.rels" in names:
            rel_root = etree.fromstring(zin.read("word/_rels/fontTable.xml.rels"))
        for rel in list(rel_root):
            if rel.get("Id") in {"rIdAlmaraiRegular", "rIdAlmaraiBold"}:
                rel_root.remove(rel)
        for rid, target in [("rIdAlmaraiRegular", "fonts/Almarai-Regular.odttf"), ("rIdAlmaraiBold", "fonts/Almarai-Bold.odttf")]:
            rel = etree.SubElement(rel_root, f"{{{ns_rel}}}Relationship")
            rel.set("Id", rid)
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
            rel.set("Target", target)
        zout.writestr("word/_rels/fontTable.xml.rels", etree.tostring(rel_root, xml_declaration=True, encoding="UTF-8", standalone="yes"))
        zout.writestr("word/fonts/Almarai-Regular.odttf", obfuscate(REGULAR_FONT.read_bytes()))
        zout.writestr("word/fonts/Almarai-Bold.odttf", obfuscate(BOLD_FONT.read_bytes()))

        ct = etree.fromstring(zin.read("[Content_Types].xml"))
        ns_ct = "http://schemas.openxmlformats.org/package/2006/content-types"
        if not ct.xpath('./ct:Default[@Extension="odttf"]', namespaces={'ct': ns_ct}):
            d = etree.SubElement(ct, f"{{{ns_ct}}}Default")
            d.set("Extension", "odttf")
            d.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
        zout.writestr("[Content_Types].xml", etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone="yes"))
    temp.replace(docx_path)


def main():
    make_assets()
    student = build_student()
    trainer = build_trainer()
    print(student)
    print(trainer)


if __name__ == "__main__":
    main()
